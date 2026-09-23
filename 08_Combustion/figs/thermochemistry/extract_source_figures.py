"""Extract the embedded source figures for Thermochemistry slide redevelopment.

These files are temporary visual references.  They are extracted from the
untracked PowerPoint so that individual visuals can be replaced deliberately
with original figures while retaining the source slide layout as a reference.
"""

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


HERE = Path(__file__).resolve().parent
SOURCE_DECK = HERE.parents[1] / "Lecture_UNTRACKED" / "Combustion_lecture_McDermott.pptx"
OUTPUT_DIRECTORY = HERE / "source"


def main() -> None:
    OUTPUT_DIRECTORY.mkdir(exist_ok=True)
    presentation = Presentation(SOURCE_DECK)

    for slide_number, slide in enumerate(presentation.slides, start=1):
        if not 6 <= slide_number <= 53:
            continue

        picture_number = 0
        media_number = 0
        for shape in slide.shapes:
            if shape.shape_type not in {
                MSO_SHAPE_TYPE.PICTURE,
                MSO_SHAPE_TYPE.MEDIA,
            }:
                continue

            if shape.shape_type == MSO_SHAPE_TYPE.MEDIA:
                media_number += 1
                visual_kind = "media"
                visual_number = media_number
            else:
                picture_number += 1
                visual_kind = "picture"
                visual_number = picture_number
            # python-pptx only creates an ImagePart for a subset of image
            # formats.  The source deck also contains TIFFs, which are still
            # ordinary package parts with a blob and can be extracted through
            # the relationship used by the picture shape.
            relationship = slide.part.rels[shape._element.blip_rId]
            source_part = relationship.target_part
            extension = Path(str(source_part.partname)).suffix.lstrip(".")
            filename = (
                f"slide-{slide_number:03d}-{visual_kind}-{visual_number:02d}.{extension}"
            )
            (OUTPUT_DIRECTORY / filename).write_bytes(source_part.blob)


if __name__ == "__main__":
    main()
